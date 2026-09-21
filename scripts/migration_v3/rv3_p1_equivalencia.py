#!/usr/bin/env python3
# ============================================================
# GAPTO MOBILE 2027
# Fichero: rv3_p1_equivalencia.py
# Ruta: scripts/migration_v3/rv3_p1_equivalencia.py
# Descripcion: RV3 / P1. Comparador versionado que decide si la Google Sheet
#              V3 actual es equivalente al snapshot lexico preservado en RUN06
#              (registros_origen_importacion.datos_origen), aplicando la regla
#              explicita de los literales legacy 141 (Migration V3 §35.1, §39;
#              expediente RV3 D-A/D-B).
#
#   Veredictos:
#     EQUIVALENTE  -> exit 0. P1 superado.
#     STOP_S3      -> exit 3. Existe al menos un delta no explicado por la
#                     regla 141 ni por equivalencia de representacion. RV3 no
#                     puede generar dataset (mandato RV3-E001 §3).
#     ERROR        -> exit 2. Entrada invalida o herramienta no ejecutada.
#
#   Garantias de la herramienta (WM D-181):
#     - Nunca escribe valores de datos V3 en la salida: solo contenedor, clave
#       origen, columna y categoria. La columna users.password no se compara
#       en claro ni se reconstruye (D-B): se registra solo si difiere.
#     - Rechaza escribir evidencia dentro del repositorio (WM §12C.11).
#     - Verifica el SHA-256 canonico del export RUN06 antes de comparar.
#     - Las claves numericas se canonicalizan (9 == 9.0) antes de emparejar;
#       sin esa regla la exportacion XLSX de Google Sheets produce falsos
#       deltas en tablas con id entero.
#
# USO:
#   python scripts/migration_v3/rv3_p1_equivalencia.py \
#       --sheet RUTA/sheet_actual.xlsx --run06 RUTA/GaptoMobile_2027_DB_RUN06_v08.xlsx \
#       --salida DIRECTORIO_FUERA_DEL_REPO
#
#   v0.2.0: SHA-256 de referencia del RUN06 re-anclado a la copia de trabajo
#           26210cb7... por decision del propietario (2026-09-21, opcion 1): el
#           fichero c13a3dd2... no se localiza; la copia produce P1 2474/26/38/1,
#           hash S1 0a180ec4... y los recuentos P5 del traspaso. El SHA retirado se
#           conserva en SHA256_RUN06_RETIRADO solo como trazabilidad (no se acepta).
# Versión: 0.2.0
# ============================================================
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import math
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

VERSION = "0.2.0"
SHA256_RUN06 = "26210cb7b950f61c7876c9a6a15bbfbd2a557475698c53ba1d056ff4c52f8556"
SHA256_RUN06_RETIRADO = "c13a3dd2104e4d58951a9d8dec100f1b4868c01de205a422df0d89a0bb965079"
HOJA_SNAPSHOT = "registros_origen_importacion"
ESPERADO = {"hojas": 44, "campos": 533, "filas_snapshot": 2501, "contenedores": 27}
SENTINEL = 141
COLUMNAS_SECRETAS = {("public.users", "password")}

# Categorias de celda
IGUAL = "IGUAL"
EQUIV_141 = "EQUIV_141"          # snapshot 141 literal, Sheet en blanco (regla §35.1)
EQUIV_REPR = "EQUIV_REPR"        # misma magnitud/texto con distinta representacion
EQUIV_NULO = "EQUIV_NULO"        # ambos vacios
DISTINTO = "DISTINTO"
SECRETO_DIFIERE = "SECRETO_DIFIERE"

# Categorias de fila
FILA_EQUIVALENTE = "EQUIVALENTE"
FILA_MODIFICADA = "MODIFICADA"
FILA_SOLO_SNAPSHOT = "SOLO_SNAPSHOT"
FILA_SOLO_SHEET = "SOLO_SHEET"
FILA_CLAVE_DUPLICADA = "CLAVE_DUPLICADA"

_T0 = time.monotonic()


def log(paso: int, total: int, msg: str) -> None:
    print(f"[{dt.datetime.now():%H:%M:%S}] [paso {paso}/{total}] {msg}", flush=True)


def latido(actual: int, total: int, etiqueta: str) -> None:
    if total and (actual == total or actual % max(1, total // 10) == 0):
        print(f"      {etiqueta}: {actual}/{total} ({100 * actual // total}%)", flush=True)


def sha256_fichero(ruta: Path) -> str:
    h = hashlib.sha256()
    with ruta.open("rb") as f:
        for bloque in iter(lambda: f.read(1 << 20), b""):
            h.update(bloque)
    return h.hexdigest()


def canon_clave(v) -> str:
    """Clave origen canonica: 9, 9.0 y '9' emparejan; el resto se compara como texto."""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and math.isfinite(v) and v.is_integer():
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    s = "" if v is None else str(v).strip()
    try:
        f = float(s)
        if math.isfinite(f) and f.is_integer() and s.replace(".", "", 1).lstrip("-").isdigit():
            return str(int(f))
    except ValueError:
        pass
    return s


def _vacio(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _es_sentinel(v) -> bool:
    return (isinstance(v, (int, float)) and not isinstance(v, bool) and v == SENTINEL) or v == str(SENTINEL)


def _num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip())
    except (TypeError, ValueError):
        return None


def _fecha_iso(v):
    if isinstance(v, dt.datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    return None


def clasificar_celda(v_snap, v_sheet) -> str:
    """Compara una celda del snapshot (JSON tipado) con la celda exportada de la Sheet."""
    if v_snap == v_sheet and type(v_snap) is type(v_sheet):
        return IGUAL
    if _es_sentinel(v_snap) and _vacio(v_sheet):
        return EQUIV_141
    if _vacio(v_snap) and _vacio(v_sheet):
        return EQUIV_NULO
    if isinstance(v_snap, bool) or isinstance(v_sheet, bool):
        if isinstance(v_snap, bool) and isinstance(v_sheet, bool):
            return IGUAL if v_snap == v_sheet else DISTINTO
        txt = {str(v_snap).strip().lower(), str(v_sheet).strip().lower()}
        return EQUIV_REPR if txt in ({"true"}, {"false"}) else DISTINTO
    iso = _fecha_iso(v_sheet)
    if iso is not None and isinstance(v_snap, str):
        s = v_snap.strip().replace("T", " ")
        if s == iso or s == iso.split(" ")[0] or (iso.endswith(" 00:00:00") and s == iso[:10]):
            return EQUIV_REPR
        if s.startswith(iso) or iso.startswith(s):
            return EQUIV_REPR
        return DISTINTO
    if v_snap is None or v_sheet is None:
        return DISTINTO
    a, b = _num(v_snap), _num(v_sheet)
    if a is not None and b is not None:
        return EQUIV_REPR if math.isclose(a, b, rel_tol=0.0, abs_tol=1e-9) else DISTINTO
    if str(v_snap).strip() == str(v_sheet).strip():
        return EQUIV_REPR
    return DISTINTO


def leer_hoja(wb, nombre):
    it = wb[nombre].iter_rows(values_only=True)
    cab = [None if x is None else str(x) for x in next(it, [])]
    filas = [r for r in it if any(v is not None and v != "" for v in r)]
    return cab, filas


def cargar_snapshot(wb):
    cab, filas = leer_hoja(wb, HOJA_SNAPSHOT)
    snap = defaultdict(dict)
    duplicadas = []
    for i, r in enumerate(filas, 1):
        reg = dict(zip(cab, r))
        cont = reg["contenedor_origen"]
        clave = canon_clave(reg["clave_origen"])
        datos = json.loads(reg["datos_origen"])
        if clave in snap[cont]:
            duplicadas.append((cont, clave))
        snap[cont][clave] = datos
        latido(i, len(filas), "registros origen")
    return snap, len(filas), duplicadas


def columna_clave(cont, snap_cont):
    clave, datos = next(iter(snap_cont.items()))
    candidatas = [c for c, v in datos.items() if canon_clave(v) == clave]
    return "id" if "id" in candidatas else (candidatas[0] if candidatas else None)


def comparar(sheet_wb, snap):
    celdas = Counter()
    filas = Counter()
    deltas = []
    inventario = {"hojas": 0, "campos": 0, "filas_sheet": 0, "hojas_con_filas": 0,
                  "hojas_vacias_con_snapshot": [], "hojas_con_filas_sin_snapshot": [],
                  "cabecera_distinta": []}
    sentinel_sheet = 0
    nombres = sheet_wb.sheetnames
    for idx, hoja in enumerate(nombres, 1):
        cab, rr = leer_hoja(sheet_wb, hoja)
        inventario["hojas"] += 1
        inventario["campos"] += len([c for c in cab if c])
        inventario["filas_sheet"] += len(rr)
        inventario["hojas_con_filas"] += 1 if rr else 0
        for r in rr:
            sentinel_sheet += sum(1 for v in r if _es_sentinel(v))
        s = snap.get(hoja)
        if not s:
            if rr:
                inventario["hojas_con_filas_sin_snapshot"].append(hoja)
                for r in rr:
                    deltas.append({"contenedor": hoja, "clave": None, "tipo": FILA_SOLO_SHEET, "columnas": []})
                    filas[FILA_SOLO_SHEET] += 1
            latido(idx, len(nombres), "hojas")
            continue
        kc = columna_clave(hoja, s)
        cab_set = {c for c in cab if c}
        cols_snap = set(next(iter(s.values())).keys())
        if cab_set != cols_snap:
            inventario["cabecera_distinta"].append(
                {"contenedor": hoja, "solo_snapshot": sorted(cols_snap - cab_set), "solo_sheet": sorted(cab_set - cols_snap)})
        sheet_rows = {}
        for r in rr:
            reg = dict(zip(cab, r))
            k = canon_clave(reg.get(kc))
            if k in sheet_rows:
                deltas.append({"contenedor": hoja, "clave": k, "tipo": FILA_CLAVE_DUPLICADA, "columnas": []})
                filas[FILA_CLAVE_DUPLICADA] += 1
            sheet_rows[k] = reg
        for k in sorted(set(s) - set(sheet_rows)):
            deltas.append({"contenedor": hoja, "clave": k, "tipo": FILA_SOLO_SNAPSHOT, "columnas": []})
            filas[FILA_SOLO_SNAPSHOT] += 1
        for k in sorted(set(sheet_rows) - set(s)):
            deltas.append({"contenedor": hoja, "clave": k, "tipo": FILA_SOLO_SHEET, "columnas": []})
            filas[FILA_SOLO_SHEET] += 1
        for k in sorted(set(s) & set(sheet_rows)):
            a, b = s[k], sheet_rows[k]
            cols_distintas = []
            for col in sorted(set(a) | set(b)):
                if (hoja, col) in COLUMNAS_SECRETAS:
                    cat = IGUAL if a.get(col) == b.get(col) else SECRETO_DIFIERE
                else:
                    cat = clasificar_celda(a.get(col), b.get(col))
                celdas[cat] += 1
                if cat in (DISTINTO, SECRETO_DIFIERE):
                    cols_distintas.append({"columna": col, "categoria": cat})
            if any(c["categoria"] == DISTINTO for c in cols_distintas):
                filas[FILA_MODIFICADA] += 1
                deltas.append({"contenedor": hoja, "clave": k, "tipo": FILA_MODIFICADA, "columnas": cols_distintas})
            else:
                filas[FILA_EQUIVALENTE] += 1
                if cols_distintas:
                    deltas.append({"contenedor": hoja, "clave": k, "tipo": "EQUIVALENTE_CON_SECRETO_REDACTADO",
                                   "columnas": cols_distintas})
        latido(idx, len(nombres), "hojas")
    return inventario, celdas, filas, deltas, sentinel_sheet


def dentro_de(ruta: Path, raiz: Path) -> bool:
    try:
        ruta.resolve().relative_to(raiz.resolve())
        return True
    except ValueError:
        return False


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="RV3 P1: equivalencia Sheet actual vs snapshot RUN06")
    ap.add_argument("--sheet", required=True, type=Path)
    ap.add_argument("--run06", required=True, type=Path)
    ap.add_argument("--salida", required=True, type=Path)
    ap.add_argument("--sha-run06", default=SHA256_RUN06)
    args = ap.parse_args(argv)
    total = 6
    try:
        import openpyxl
    except ImportError:
        print("ERROR: falta openpyxl", file=sys.stderr)
        return 2
    repo = Path(__file__).resolve().parents[2]
    if dentro_de(args.salida, repo):
        print("ERROR: la salida debe quedar FUERA del repositorio (WM §12C.11)", file=sys.stderr)
        return 2
    args.salida.mkdir(parents=True, exist_ok=True)

    log(1, total, "Verificando identidad de las entradas")
    sha_run06 = sha256_fichero(args.run06)
    sha_sheet = sha256_fichero(args.sheet)
    if sha_run06 != args.sha_run06:
        print(f"STOP S3: SHA-256 de RUN06 no canonico ({sha_run06})")
        return 3

    log(2, total, "Cargando snapshot lexico de RUN06")
    wb_r = openpyxl.load_workbook(args.run06, read_only=True, data_only=True)
    snap, n_snap, dup_snap = cargar_snapshot(wb_r)
    sentinel_snap = sum(1 for m in snap.values() for d in m.values() for v in d.values() if _es_sentinel(v))
    none_snap = sum(1 for m in snap.values() for d in m.values() for v in d.values() if v in ("None", "NONE"))

    log(3, total, "Cargando Google Sheet exportada")
    wb_s = openpyxl.load_workbook(args.sheet, read_only=True, data_only=True)

    log(4, total, "Comparando hoja a hoja, clave a clave y celda a celda")
    inv, celdas, filas, deltas, sentinel_sheet = comparar(wb_s, snap)

    log(5, total, "Evaluando veredicto")
    estructurales = []
    for k, esperado in (("hojas", ESPERADO["hojas"]), ("campos", ESPERADO["campos"])):
        if inv[k] != esperado:
            estructurales.append(f"{k}: esperado {esperado}, Sheet {inv[k]}")
    if n_snap != ESPERADO["filas_snapshot"]:
        estructurales.append(f"filas snapshot: esperado {ESPERADO['filas_snapshot']}, RUN06 {n_snap}")
    if len(snap) != ESPERADO["contenedores"]:
        estructurales.append(f"contenedores snapshot: esperado {ESPERADO['contenedores']}, RUN06 {len(snap)}")
    if dup_snap:
        estructurales.append(f"claves duplicadas en snapshot: {len(dup_snap)}")
    if inv["cabecera_distinta"]:
        estructurales.append(f"cabeceras distintas: {len(inv['cabecera_distinta'])}")
    no_explicados = [d for d in deltas if d["tipo"] in
                     (FILA_MODIFICADA, FILA_SOLO_SHEET, FILA_SOLO_SNAPSHOT, FILA_CLAVE_DUPLICADA)]
    veredicto = "EQUIVALENTE" if not (estructurales or no_explicados) else "STOP_S3"

    por_contenedor = defaultdict(Counter)
    for d in no_explicados:
        por_contenedor[d["contenedor"]][d["tipo"]] += 1
    columnas_modificadas = defaultdict(Counter)
    for d in no_explicados:
        for c in d["columnas"]:
            if c["categoria"] == DISTINTO:
                columnas_modificadas[d["contenedor"]][c["columna"]] += 1

    informe = {
        "herramienta": "scripts/migration_v3/rv3_p1_equivalencia.py",
        "version": VERSION,
        "ejecutado_utc": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "veredicto": veredicto,
        "entradas": {"run06_sha256": sha_run06, "run06_sha256_canonico": sha_run06 == SHA256_RUN06,
                     "sheet_export_sha256": sha_sheet,
                     "nota_sheet": "El SHA del export XLSX de Google Sheets identifica ESTA descarga; no es identidad estable de la Sheet."},
        "inventario_sheet": {k: v for k, v in inv.items()},
        "snapshot": {"filas": n_snap, "contenedores": len(snap), "claves_duplicadas": len(dup_snap),
                     "literales_141": sentinel_snap, "literales_None_NONE": none_snap},
        "sheet_literales_141_no_vacios": sentinel_sheet,
        "celdas": dict(celdas),
        "filas": dict(filas),
        "problemas_estructurales": estructurales,
        "deltas_no_explicados_por_contenedor": {k: dict(v) for k, v in sorted(por_contenedor.items())},
        "columnas_modificadas_por_contenedor": {k: dict(v) for k, v in sorted(columnas_modificadas.items())},
        "deltas": deltas,
    }
    marca = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    salida = args.salida / f"rv3_p1_equivalencia_{marca}.json"
    salida.write_text(json.dumps(informe, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")

    log(6, total, f"Veredicto: {veredicto}")
    print(f"      filas: {dict(filas)}")
    print(f"      celdas: {dict(celdas)}")
    for e in estructurales:
        print(f"      ESTRUCTURAL: {e}")
    for c, v in sorted(por_contenedor.items()):
        print(f"      DELTA {c}: {dict(v)}")
    print(f"      informe: {salida.name}  sha256={sha256_fichero(salida)}")
    print(f"      duracion: {time.monotonic() - _T0:.1f} s")
    return 0 if veredicto == "EQUIVALENTE" else 3


if __name__ == "__main__":
    sys.exit(main())
