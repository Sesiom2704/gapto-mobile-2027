# ============================================================
# GAPTO MOBILE 2027
# Fichero: rv3_fuente.py
# Ruta: scripts/migration_v3/rv3_fuente.py
# Descripcion: RV3 / P2-P4. Libreria unica de fuente para el pipeline RV3:
#
#   - B0: carga del snapshot origen preservado en RUN06
#     (registros_origen_importacion). Para cada fila conserva la cadena JSON
#     EXACTA de la celda datos_origen como SERIALIZACION_SNAPSHOT_RUN06 y
#     calcula sha256_registro con el algoritmo RV3_SOURCE_JSONCELL_V1
#     (SHA-256 de los bytes UTF-8 exactos de esa cadena). NO es el algoritmo
#     historico de F01 ni texto original de la BD V3 (RV3-D002-B).
#   - S1: carga del estado semantico congelado de la Google Sheet (P1b).
#   - Normalizacion P3: literal + valor + estado. Regla contextual del 141
#     (RV3-D002-C): 141 es sentinel de ausencia salvo que dominio y contexto
#     prueben valor real; unico caso contractual: prestamo_cuota.num_cuota
#     dentro de la secuencia 1..N de su prestamo. None/NONE son marcadores de
#     ausencia legacy (Migration V3 §2). El literal se conserva siempre.
#   - Identidad P4: UUIDv5 con namespace canonico y nombre
#     v3|{contenedor}|{clave}|{tabla_destino}|{rol_destino}.
#
#   Una sola implementacion sirve a B0 y S1 (mandato RV3-E001-R1 §5): S1 solo
#   cambia el modo de evidencia, nunca las reglas.
#
# Versión: 0.1.0
# ============================================================
from __future__ import annotations

import hashlib
import importlib.util
import json
import uuid
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

VERSION = "0.1.0"
NAMESPACE_V3 = uuid.UUID("0e687b90-e2c6-5bf2-8ac1-99eecfe4f811")
ALGORITMO_HASH_REGISTRO = "RV3_SOURCE_JSONCELL_V1"
SEMANTICA_TEXTO = "SERIALIZACION_SNAPSHOT_RUN06"
SENTINEL_141 = 141
MARCADORES_AUSENCIA_TEXTO = frozenset({"None", "NONE"})
CONTEXTO_141_REAL = {("public.prestamo_cuota", "num_cuota")}

# Estados de normalizacion
CONOCIDO = "CONOCIDO"
NULO = "NULO"                            # vacio/None real de la fuente
AUSENCIA_141 = "AUSENCIA_LEGACY_141"     # sentinel 141 -> desconocido
AUSENCIA_NONE = "AUSENCIA_LEGACY_NONE"   # literal None/NONE -> desconocido
REAL_141 = "VALOR_REAL_141"              # 141 probado como valor

_AQUI = Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location("rv3_p1_equivalencia", _AQUI / "rv3_p1_equivalencia.py")
p1 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(p1)


# ------------------------------------------------------------------ identidad
def uuid_v3(contenedor: str, clave: str, tabla: str, rol: str) -> str:
    for parte, nombre in ((contenedor, "contenedor"), (clave, "clave"), (tabla, "tabla"), (rol, "rol")):
        if parte is None or str(parte) == "" or "|" in str(parte) and nombre != "clave":
            raise ValueError(f"identidad invalida: {nombre}={parte!r}")
    return str(uuid.uuid5(NAMESPACE_V3, f"v3|{contenedor}|{clave}|{tabla}|{rol}"))


def uuid_origen(contenedor: str, clave: str) -> str:
    return uuid_v3(contenedor, clave, "registros_origen_importacion", "origen")


# ------------------------------------------------------------------ B0
@dataclass(frozen=True)
class RegistroB0:
    contenedor: str
    clave: str
    datos: dict
    texto: str                 # SERIALIZACION_SNAPSHOT_RUN06, byte-exacta
    sha256_registro: str       # RV3_SOURCE_JSONCELL_V1
    id_run06: str
    numero_fila_origen: int | None


def hash_registro(texto: str) -> str:
    return hashlib.sha256(texto.encode("utf-8")).hexdigest()


def cargar_b0(ruta_run06: Path, sha_esperado: str = p1.SHA256_RUN06) -> dict[str, dict[str, RegistroB0]]:
    import openpyxl
    ruta_run06 = Path(ruta_run06)
    sha = p1.sha256_fichero(ruta_run06)
    if sha != sha_esperado:
        raise ValueError(f"RUN06 no canonico: {sha}")
    wb = openpyxl.load_workbook(ruta_run06, read_only=True, data_only=True)
    cab, filas = p1.leer_hoja(wb, p1.HOJA_SNAPSHOT)
    b0: dict[str, dict[str, RegistroB0]] = defaultdict(dict)
    for r in filas:
        reg = dict(zip(cab, r))
        texto = reg["datos_origen"]
        if not isinstance(texto, str):
            raise ValueError("datos_origen no es cadena JSON")
        cont = reg["contenedor_origen"]
        clave = p1.canon_clave(reg["clave_origen"])
        if clave in b0[cont]:
            raise ValueError(f"clave duplicada en B0: {cont}/{clave}")
        nf = reg.get("numero_fila_origen")
        b0[cont][clave] = RegistroB0(cont, clave, json.loads(texto), texto, hash_registro(texto),
                                     reg["id"], int(nf) if isinstance(nf, (int, float)) else None)
    return dict(b0)


# ------------------------------------------------------------------ S1
def cargar_s1(ruta_sheet: Path) -> dict[str, dict]:
    """Estado semantico S1 (mismas reglas de clave que P1b); valores tal como exporta la Sheet."""
    import openpyxl
    spec = importlib.util.spec_from_file_location("rv3_p1b_manifest_s1", _AQUI / "rv3_p1b_manifest_s1.py")
    p1b = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(p1b)
    wb = openpyxl.load_workbook(ruta_sheet, read_only=True, data_only=True)
    s1 = {}
    for hoja in wb.sheetnames:
        cab, filas = p1.leer_hoja(wb, hoja)
        if not filas:
            s1[hoja] = {"campos": [c for c in cab if c], "filas": {}}
            continue
        kc = p1b.columna_clave(hoja, cab)
        regs = {}
        for r in filas:
            d = dict(zip(cab, r))
            regs[p1.canon_clave(d.get(kc))] = d
        s1[hoja] = {"campos": [c for c in cab if c], "filas": regs}
    return s1


# ------------------------------------------------------------------ normalizacion
@dataclass(frozen=True)
class Celda:
    literal: object
    valor: object
    estado: str


def secuencias_cuotas(filas_cuota: dict) -> dict:
    """prestamo_id -> lista de num_cuota tal como vienen (literal)."""
    seq = defaultdict(list)
    for d in filas_cuota.values():
        seq[str(d.get("prestamo_id"))].append(d.get("num_cuota"))
    return seq


def _141_es_real(contenedor: str, columna: str, fila: dict, contexto: dict) -> bool:
    if (contenedor, columna) not in CONTEXTO_141_REAL:
        return False
    # prestamo_cuota.num_cuota: 141 es real si la secuencia del prestamo es
    # exactamente 1..N (sin huecos ni duplicados) y N >= 141.
    seq = contexto.get("cuotas", {}).get(str(fila.get("prestamo_id")))
    if not seq:
        return False
    try:
        nums = sorted(int(float(x)) for x in seq)
    except (TypeError, ValueError):
        return False
    return nums == list(range(1, len(nums) + 1)) and len(nums) >= SENTINEL_141


def normalizar(contenedor: str, columna: str, v, fila: dict, contexto: dict) -> Celda:
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return Celda(v, None, NULO)
    if isinstance(v, str) and v in MARCADORES_AUSENCIA_TEXTO:
        return Celda(v, None, AUSENCIA_NONE)
    if p1._es_sentinel(v):
        if _141_es_real(contenedor, columna, fila, contexto):
            return Celda(v, SENTINEL_141, REAL_141)
        return Celda(v, None, AUSENCIA_141)
    return Celda(v, v, CONOCIDO)


def contexto_de(fuente: dict[str, dict]) -> dict:
    """fuente: contenedor -> {clave: dict_fila}. Construye el contexto de normalizacion."""
    return {"cuotas": secuencias_cuotas(fuente.get("public.prestamo_cuota", {}))}


def normalizar_fuente(fuente: dict[str, dict]) -> dict[str, dict[str, dict[str, Celda]]]:
    ctx = contexto_de(fuente)
    return {c: {k: {col: normalizar(c, col, v, fila, ctx) for col, v in fila.items()}
                for k, fila in filas.items()}
            for c, filas in fuente.items()}


def fuente_b0(b0: dict[str, dict[str, RegistroB0]]) -> dict[str, dict]:
    return {c: {k: r.datos for k, r in regs.items()} for c, regs in b0.items()}


def fuente_s1(s1: dict[str, dict]) -> dict[str, dict]:
    return {h: e["filas"] for h, e in s1.items() if e["filas"]}
